"""Generate deliverables/summary.docx (2-3 page project summary). Run: python -I -X utf8 deliverables/make_summary.py"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "summary.docx"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2)
    s.left_margin = s.right_margin = Cm(2.2)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(4)
for name, size in (("Title", 24), ("Heading 1", 13)):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = ACCENT
doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(3)


def h(text):
    doc.add_heading(text, level=1)


def p(text, bold_lead=None):
    para = doc.add_paragraph()
    if bold_lead:
        para.add_run(bold_lead).bold = True
    para.add_run(text)
    return para


def bullets(items, style="List Bullet"):
    for item in items:
        lead, _, rest = item.partition("|")
        para = doc.add_paragraph(style=style)
        if rest:
            para.add_run(lead).bold = True
            para.add_run(rest)
        else:
            para.add_run(lead)
        para.paragraph_format.space_after = Pt(2)


doc.add_heading("Treaty Room", level=0)
sub = doc.add_paragraph()
sub.add_run("A sealed research room that no one owns").italic = True
meta = doc.add_paragraph()
meta.add_run("Team: ").bold = True
meta.add_run("Sunway PowerRanger (judged in Malaysia)    ")
meta.add_run("Event: ").bold = True
meta.add_run("RMIT x Tide UNLOCKED, Challenge 2\n")
meta.add_run("Repository: ").bold = True
meta.add_run("https://github.com/ChewEzekiel/treaty-room    ")
meta.add_run("Demo video: ").bold = True
meta.add_run("demo/video.mp4 (recorded from branch walkthrough-demo with simulated Tide calls, see below)")

h("Problem")
p(
    "Most research never sees daylight. Failed trials, negative results and early assay data sit in drawers at "
    "universities, drug companies and hospitals, so other teams repeat the same dead ends and safety signals surface "
    "late. The groups that hold this evidence are rivals: they compete for grants, patents and first publication. "
    "Pooling unpublished results is in everyone's interest, but nobody wants to be the one who shares first."
)

h("The trust constraint")
p(
    "Today, pooling requires trusting someone. A hosted data room means trusting the host, who can read, leak or be "
    "breached. A lead partner running the room can read everyone else's results first and publish first. Data trusts, "
    "escrow agents and confidentiality agreements move that trust to a third party; they do not remove it, and a "
    "single administrator with a master key can still open the room early. That is why the room is rarely built."
)
p(
    "When no one, including the host and any single institution, can read a deposit before the agreed moment, "
    "competitors can pool evidence before publication without giving up their lead.",
    "What changes when the constraint is removed: ",
)

h("Solution")
p(
    "Treaty Room is a sealed room for a consortium. Each member's deposit is encrypted in their own browser to the "
    "room. The host's database holds only ciphertext and metadata. The room is sealed: even depositors cannot open "
    "it. It unseals for every member at the same moment, and only when 2 of 3 institution admins co-sign one "
    "TideCloak change request: the \"treaty\". A rogue steward who proposes early access stalls at 1 of 2 "
    "signatures. Joining a room is likewise a governed change request."
)
p("The prototype runs two synthetic rooms:")
bullets([
    "Alzheimer's blood-test consortium|: can a p-tau217 blood test replace PET scans? Northbridge University, "
    "Helix Therapeutics and St Aurelia Hospital.",
    "Dengue vaccine safety signals|: do Phase II sites in Southeast Asia show the same early signal? Sunway "
    "Institute, Helix Therapeutics and KL General Hospital.",
])

h("Why TideCloak matters")
p("TideCloak is what makes the seal real rather than a promise in our code:")
bullets([
    "Login via Tide.| We store no passwords. Tokens are threshold-signed by Tide's ORK network, so our server can "
    "verify them but cannot mint them.",
    "Two network-enforced policies per room.| A realm admin signs both once, in their Tide enclave (\"1. Sign "
    "deposit policy\" and \"2. Sign reading policy\"), under Tide's built-in RelationshipAttestation:1 contract with "
    "the room id as the only tag. Deposit policy: subject = group /rooms/<id>, role _tide_encryptor. Reading "
    "policy: subject = realm role room-<id>-reader, role _tide_decryptor. On every encrypt and decrypt, Tide's "
    "decentralised nodes check the user's doken against the policy. Plaintext never exists on our server.",
    "Quorum-governed seal.| Members get _tide_encryptor through the room group. Nobody holds room-<id>-reader at "
    "first. The unseal is one IGA change request (\"Propose unseal: <room>\") that maps both _tide_decryptor and "
    "room-<id>-reader onto the room group, under Quorum-Enforced Governance (Tide mode). Because the reader role is "
    "per room, opening one room opens no other. Approvals are enclave signatures from each admin's own browser, "
    "with threshold max(1, floor(0.7 x admins)), which is 2 of 3. The server-side \"authorize\" call is refused in "
    "multi-admin Tide mode, so the host cannot approve on anyone's behalf.",
    "Visible proof.| The Operator view shows everything the host or a database thief holds, and \"Try to decrypt "
    "with the stolen data\" shows the Tide network's refusal. \"Force as operator\" (realm admins, only with "
    "DEMO_RED_TEAM=1) shows TideCloak refusing a master-admin approval.",
])
p(
    "With ordinary login and backend permissions, the host's administrator could flip a database flag and read "
    "everything. Here the host never holds a key and cannot fake an approval.",
    "Without TideCloak: ",
)

h("Demo flow")
bullets([
    "Deposit.| Admins from the three institutions sign in, click \"Load my synthetic example\" and \"Encrypt & "
    "deposit\". The server stores ciphertext only.",
    "Sealed.| \"Open\" / \"Open all\" is refused by the Tide network: no one holds the room's reader role yet, "
    "depositors included.",
    "Rogue steward.| One admin files \"Propose unseal\" on Governance and signs it; it waits at 1 of 2. \"Force as "
    "operator\" tries the server-side authorize lane with the master-admin password; TideCloak refuses.",
    "Treaty.| A second admin signs in their own enclave. Quorum is reached and the change commits, mapping "
    "_tide_decryptor and room-<id>-reader onto that room's group only.",
    "Unseal.| Members click \"Refresh access\" for a fresh token, then \"Open all\": the room opens for all of them.",
    "Operator view.| The raw database file (ciphertext, signed policies, metadata) and the network's refusal to "
    "decrypt it.",
], style="List Number")

h("Tide login and the demo video")
p(
    "The app on main requires Tide login. There is no simulated sign-in, enclave or encryption: users sign in through "
    "TideCloak with a Tide account, admins sign policies and change requests in their own Tide enclave, and every "
    "encrypt and decrypt is checked by Tide's ORK nodes.",
    "Main is Tide only. ",
)
p(
    "On 22 October, Tide's public ORK network had nodes down (ork2, ork3 and ork7 returned HTTP 503), and creating a "
    "new Tide account needs every selected ORK. After ork2 recovered, new accounts were refused by the free-tier cap "
    "(TIDE-ORK-VOUCHER-SUBSCRIPTION_EXCEEDED_CONCURRENT), so nobody could enrol or complete a login. To still show "
    "the screens and flows, the demo video was recorded from a separate branch, walkthrough-demo, where the Tide SDK "
    "calls (sign-in, enclave co-signing, encryption) are simulated with synthetic personas and a banner on every page "
    "says so. The video does not show live Tide, and the simulation is not part of main.",
    "Honest disclosure: ",
)
p(
    "TideCloak on the official image in a GitHub Codespace; realm consortium licensed with a Tide vendor key; IGA in "
    "Tide mode, so governed admin writes are answered with HTTP 202 change requests (266 approved, including the two "
    "per-room reader roles); and the Tide adapter config exported for the app. The read-back realm state, logs and "
    "outage timeline are in docs/evidence/live-tidecloak.md.",
    "What ran live: ",
)

h("Impact")
p(
    "Researchers can share dead ends and early signals without losing priority, because no partner reads first. "
    "Funders and ethics boards see fewer duplicated cohorts. Patients and trial volunteers benefit when a safety "
    "signal seen at one site is pooled with the others while the trial is still running. The same pattern fits any "
    "rival group that must pool sensitive evidence under an agreed rule: hospital incident data, cross-bank fraud "
    "signals, or sealed-bid evaluations."
)

h("Limitations")
bullets([
    "No Tide account could be created on 22 October, so Tide login, enclave signing of the room policies and "
    "RelationshipAttestation:1 encrypt and decrypt were not run live. The video shows them simulated.",
    "Main has no fallback: using the app needs an enrolled Tide account.",
    "Branch walkthrough-demo is a recording build with no cryptographic enforcement; it must never be deployed.",
    "The institution in a user's profile is self-asserted; access decisions use the signed groups and roles.",
    "Metadata (author name and institution, room, post type, time and size) is visible to the host.",
    "Whoever serves the JavaScript is trusted: a malicious host could ship code that exfiltrates plaintext after "
    "decryption in the browser.",
    "Revocation is forward-only; an admitted member can copy plaintext they have read.",
    "The host can delete or withhold ciphertext (denial of service), though it cannot read or forge it.",
    "Two colluding institutions can unseal over the third: that is the agreed treaty rule.",
    "The demo uses synthetic data and the TideCloak development image (tideorg/tidecloak-dev).",
])

h("Running it (summary)")
p("Full steps and troubleshooting are in RUN_INSTRUCTIONS.md (also supplied as a separate deliverable).")
bullets([
    "TideCloak:| it must run on Linux with Docker (Tide's Windows native library is a stub); we used a GitHub "
    "Codespace. Copy app/.env.example to app/.env and set KC_BOOTSTRAP_ADMIN_PASSWORD, TIDE_LICENSE_EMAIL and "
    "TIDE_TERMS_ACCEPTED=yes (the last only after reading https://tide.org/legal), run bash "
    "scripts/codespace-tidecloak.sh and forward ports with gh codespace ports forward 8080:8080 3000:3000.",
    "App:| enrol the first admin from the printed invite link, then cd app && npm install && npm run dev and open "
    "http://localhost:3000. Sign-in always goes through Tide.",
    "Enrol three admins (one per institution) for a real 2-of-3 quorum; an admin signs each room's two policies once.",
    "Our venue network intercepted TLS, so npm used --registry=https://registry.yarnpkg.com and the browser reached "
    "Tide through a SOCKS proxy via the Codespace.",
])

foot = doc.add_paragraph()
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = foot.add_run("Built on Tide's nextjs-e2ee-vault template. Interface: Geist, zinc with an emerald accent, light "
                 "and dark themes, Phosphor icons. All institutions, people and results are synthetic.")
r.italic = True
r.font.size = Pt(9)

doc.save(OUT)
print(f"wrote {OUT}")
